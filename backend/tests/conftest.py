import fitz
import pytest
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook


def _text_image(text="Hello OpenConvert 12345", size=(900, 200)):
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 48)
    except OSError:
        font = ImageFont.load_default()
    d.text((30, 60), text, fill="black", font=font)
    return img


@pytest.fixture
def image_path(tmp_path):
    p = tmp_path / "scan.png"
    _text_image().save(p)
    return str(p)


@pytest.fixture
def text_pdf(tmp_path):
    p = tmp_path / "text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Native text layer document for OpenConvert testing.", fontsize=14)
    doc.save(p)
    return str(p)


@pytest.fixture
def scanned_pdf(tmp_path, image_path):
    p = tmp_path / "scanned.pdf"
    doc = fitz.open()
    for _ in range(2):
        page = doc.new_page(width=612, height=200)
        page.insert_image(page.rect, filename=image_path)
    doc.save(p)
    return str(p)


@pytest.fixture
def docx_path(tmp_path):
    p = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("OpenConvert Word document")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Name"
    table.rows[0].cells[1].text = "Value"
    document.save(p)
    return str(p)


@pytest.fixture
def xlsx_path(tmp_path):
    p = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Results"
    sheet.append(["Name", "Value"])
    sheet.append(["OpenConvert", 123])
    workbook.save(p)
    return str(p)
