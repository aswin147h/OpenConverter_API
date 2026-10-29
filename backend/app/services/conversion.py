"""Route adapters. Each takes (src, dst, lang, progress) and writes dst."""
import csv
import os
import shutil
import subprocess
import tempfile
import uuid
from pathlib import Path

from docx import Document
from openpyxl import load_workbook

from . import ocr, pdf
from ..errors import ConversionError


def image_to_txt(src, dst, lang="eng", progress=None):
    text = ocr.image_to_text(ocr.open_image(src), lang)
    with open(dst, "w", encoding="utf-8") as f:
        f.write(text)
    if progress:
        progress(1, 1)


def image_to_searchable_pdf(src, dst, lang="eng", progress=None):
    data = ocr.image_to_pdf_bytes(ocr.open_image(src), lang)
    with open(dst, "wb") as f:
        f.write(data)
    if progress:
        progress(1, 1)


def pdf_to_txt(src, dst, lang="eng", progress=None):
    text = pdf.pdf_to_text(src, lang, progress=progress)
    with open(dst, "w", encoding="utf-8") as f:
        f.write(text)


def pdf_to_searchable_pdf(src, dst, lang="eng", progress=None):
    pdf.pdf_to_searchable_pdf(src, dst, lang, progress=progress)


def docx_to_txt(src, dst, lang="eng", progress=None):
    document = Document(src)
    lines = []
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            lines.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            lines.append("\t".join(cell.text for cell in row.cells))
    with open(dst, "w", encoding="utf-8", newline="") as output:
        output.write("\n".join(lines))
    if progress:
        progress(1, 1)


def xlsx_to_csv(src, dst, lang="eng", progress=None):
    workbook = load_workbook(src, read_only=True, data_only=True)
    worksheets = workbook.worksheets
    with open(dst, "w", encoding="utf-8", newline="") as output:
        writer = csv.writer(output)
        for worksheet_index, worksheet in enumerate(worksheets, start=1):
            if worksheet_index > 1:
                writer.writerow([])
            writer.writerow([f"Worksheet: {worksheet.title}"])
            for row in worksheet.iter_rows(values_only=True):
                writer.writerow(["" if value is None else value for value in row])
    workbook.close()
    if progress:
        progress(1, 1)


def _libreoffice_executable():
    candidates = [
        shutil.which("soffice"),
        shutil.which("libreoffice"),
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/usr/bin/libreoffice",
    ]
    for candidate in candidates:
        if candidate and os.path.isfile(candidate):
            return candidate
    raise ConversionError("office_engine_unavailable", "LibreOffice is not installed.")


def office_to_pdf(src, dst, lang="eng", progress=None):
    src_path = Path(src)
    dst_path = Path(dst)
    executable = _libreoffice_executable()
    profile = Path(tempfile.gettempdir()) / "openconvert-office" / uuid.uuid4().hex
    profile.mkdir(parents=True, exist_ok=True)

    try:
        with tempfile.TemporaryDirectory(prefix="openconvert-office-") as work_dir:
            command = [
                executable,
                f"-env:UserInstallation=file:///{profile.as_posix()}",
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                work_dir,
                str(src_path.resolve()),
            ]
            kwargs = {"timeout": 120}
            if os.name == "nt":
                kwargs["creationflags"] = subprocess.CREATE_NEW_PROCESS_GROUP
            result = subprocess.run(
                command,
                capture_output=True,
                text=True,
                check=False,
                **kwargs,
            )
            if result.returncode != 0:
                raise ConversionError("office_conversion_failed", "Office conversion failed.")
            generated = Path(work_dir) / f"{src_path.stem}.pdf"
            if not generated.is_file():
                raise ConversionError("office_conversion_failed", "LibreOffice produced no PDF.")
            shutil.copyfile(generated, dst_path)
    except subprocess.TimeoutExpired as exc:
        raise ConversionError("office_timeout", "Office conversion timed out.") from exc
    finally:
        shutil.rmtree(profile, ignore_errors=True)

    if progress:
        progress(1, 1)


def docx_to_pdf(src, dst, lang="eng", progress=None):
    office_to_pdf(src, dst, lang=lang, progress=progress)


def xlsx_to_pdf(src, dst, lang="eng", progress=None):
    office_to_pdf(src, dst, lang=lang, progress=progress)


def pptx_to_pdf(src, dst, lang="eng", progress=None):
    office_to_pdf(src, dst, lang=lang, progress=progress)


def pdf_to_docx(src, dst, lang="eng", progress=None):
    doc = pdf.open_pdf(src)
    try:
        total = doc.page_count
        document = Document()
        for index, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                for paragraph in text.split("\n\n"):
                    document.add_paragraph(paragraph)
            else:
                document.add_paragraph("[Text could not be extracted from this page.]")
            if progress:
                progress(index, total)
        document.save(dst)
    finally:
        doc.close()
